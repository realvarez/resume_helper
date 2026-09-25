import io
from unittest.mock import MagicMock, patch
import docx
import pytest
from fastapi import HTTPException, UploadFile
from app.services.parser import _decode_text, build_extra_context, extract_text


def _make_upload(filename: str, content: bytes) -> UploadFile:
    return UploadFile(filename=filename, file=io.BytesIO(content))


def test_decode_text_utf8_and_latin1():
    assert _decode_text(b"Hello world") == "Hello world"
    latin1_bytes = "Café".encode("latin-1")
    assert _decode_text(latin1_bytes) == "Café"


def test_extract_text_empty_upload():
    assert extract_text(None) == ""
    upload = UploadFile(filename="", file=io.BytesIO(b""))
    assert extract_text(upload) == ""


def test_extract_text_unsupported_extension():
    upload = _make_upload("malicious.exe", b"binary")
    with pytest.raises(HTTPException) as exc:
        extract_text(upload)
    assert exc.value.status_code == 400
    assert "Unsupported file type '.exe'" in exc.value.detail


def test_extract_text_max_upload_size():
    upload = _make_upload("large.txt", b"A" * (6 * 1024 * 1024))
    with pytest.raises(HTTPException) as exc:
        extract_text(upload)
    assert exc.value.status_code == 413
    assert "too large" in exc.value.detail


def test_extract_text_empty_file_min_chars():
    upload = _make_upload("empty.txt", b"")
    assert extract_text(upload, min_chars=0) == ""

    upload2 = _make_upload("empty.txt", b"")
    with pytest.raises(HTTPException) as exc:
        extract_text(upload2, min_chars=50)
    assert exc.value.status_code == 400
    assert "Uploaded file is empty" in exc.value.detail


def test_extract_text_min_chars_threshold():
    upload = _make_upload("short.txt", b"Too short")
    with pytest.raises(HTTPException) as exc:
        extract_text(upload, min_chars=50)
    assert exc.value.status_code == 422
    assert "Could not extract enough readable text" in exc.value.detail


def test_extract_text_txt_and_md():
    txt_upload = _make_upload("resume.txt", b"Senior Python Developer with 10 years experience.")
    assert extract_text(txt_upload) == "Senior Python Developer with 10 years experience."

    md_upload = _make_upload("resume.md", b"# Jane Doe\nSoftware Engineer")
    assert extract_text(md_upload) == "# Jane Doe\nSoftware Engineer"


def test_extract_text_docx():
    doc = docx.Document()
    doc.add_paragraph("Paragraph 1: Summary")
    doc.add_paragraph("Paragraph 2: Experience")
    stream = io.BytesIO()
    doc.save(stream)

    upload = _make_upload("resume.docx", stream.getvalue())
    extracted = extract_text(upload)
    assert "Paragraph 1: Summary" in extracted
    assert "Paragraph 2: Experience" in extracted


def test_extract_text_pdf():
    # Mock pdfplumber to avoid binary font/pdf dependency
    mock_page1 = MagicMock()
    mock_page1.extract_text.return_value = "Page 1 Content"
    mock_page2 = MagicMock()
    mock_page2.extract_text.return_value = "Page 2 Content"

    mock_pdf = MagicMock()
    mock_pdf.pages = [mock_page1, mock_page2]
    mock_pdf.__enter__.return_value = mock_pdf

    with patch("pdfplumber.open", return_value=mock_pdf):
        upload = _make_upload("resume.pdf", b"%PDF-1.4 dummy bytes")
        extracted = extract_text(upload)
        assert extracted == "Page 1 Content\nPage 2 Content"


def test_build_extra_context():
    f1 = _make_upload("cert.txt", b"AWS Solutions Architect Certified 2023")
    f2 = _make_upload("awards.md", b"Top Engineer Q3 2022")
    notes = "Looking for remote roles in Canada only."

    context = build_extra_context(supporting_files=[f1, f2], extra_notes=notes)
    assert "### Source: cert.txt\nAWS Solutions Architect Certified 2023" in context
    assert "### Source: awards.md\nTop Engineer Q3 2022" in context
    assert "### Candidate Extra Notes & Highlights\nLooking for remote roles in Canada only." in context


def test_build_extra_context_empty():
    assert build_extra_context([], "") == ""
    assert build_extra_context(None, None) == ""
